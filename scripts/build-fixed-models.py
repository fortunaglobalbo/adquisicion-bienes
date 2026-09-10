"""Build the fixed Word models from retained references; never overwrite originals."""
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import json, hashlib

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / 'PROCESO DE ADQUISICION'
OUT = ROOT / 'templates' / 'ende'
OUT.mkdir(parents=True, exist_ok=True)
catalog = []

def token(key): return '{{' + key + '}}'
def word_field(p, instruction):
    f=OxmlElement('w:fldSimple'); f.set(qn('w:instr'),instruction)
    r=OxmlElement('w:r'); t=OxmlElement('w:t'); t.text=''; r.append(t); f.append(r); p._p.append(f)
def replace_p(p, value):
    runs = p.runs
    if runs:
        runs[0].text = value
        for r in runs[1:]: r.text = ''
    else: p.add_run(value)
def field(key, label, normative=False, hint=''):
    return dict(key=key, label=label, normative=normative, hint=hint)
def save(doc, number, title, fields, columns=None, source=''):
    file = f'{number:02d}.docx'
    doc.core_properties.author = 'ENDE Deoruro S.A.'
    doc.core_properties.last_modified_by = 'ENDE Deoruro S.A.'
    doc.core_properties.title = title
    doc.core_properties.comments = ''
    doc.save(OUT / file)
    catalog.append(dict(number=number, title=title, file=file, version='2026-09-10.1', fields=fields, columns=columns or [], source=source,
                        sha256=hashlib.sha256((OUT/file).read_bytes()).hexdigest()))
def base(title, subtitle=''):
    d = Document()
    sec = d.sections[0]; sec.page_width=Cm(21); sec.page_height=Cm(29.7)
    sec.top_margin=Cm(1.6); sec.bottom_margin=Cm(1.6); sec.left_margin=Cm(1.8); sec.right_margin=Cm(1.8)
    style=d.styles['Normal']; style.font.name='Arial'; style.font.size=Pt(10)
    style.paragraph_format.space_after=Pt(5)
    p=d.add_paragraph(); p.add_run().add_picture(str(ROOT/'public/logo-ende-deoruro.png'),width=Cm(3))
    p=d.add_paragraph('DISTRIBUIDORA DE ELECTRICIDAD ENDE DEORURO S.A.'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].bold=True
    p=d.add_paragraph(title); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.runs[0].bold=True
    if subtitle:
        p=d.add_paragraph(subtitle); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    return d
def para(d, label, key): d.add_paragraph(label + token(key))
def table(d, headers, keys):
    t=d.add_table(rows=2, cols=len(headers)); t.style='Table Grid'
    for c,v in zip(t.rows[0].cells,headers): c.text=v
    for c,k in zip(t.rows[1].cells,keys): c.text=token('items.'+k)
    prop=OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(prop)
    return [dict(key=k,label=h) for h,k in zip(headers,keys)]
def signatures(d, labels):
    d.add_paragraph('')
    t=d.add_table(rows=1,cols=len(labels))
    for c,label in zip(t.rows[0].cells,labels): c.text='\n________________________\n'+label+'\nFirma y sello'

# Original TDR: preserve page furniture, index and section layout, replace all purchase content.
p=SRC/'1TDR/TDR_ENDE.docx'; d=Document(p)
fields=[field('objeto','Objeto de la compra'),field('resumen','Resumen de ítems'),field('mes_anio','Mes y año'),
        field('elaborado','Elaborado por'),field('revisado','Revisado por'),field('aprobado','Aprobado por')]
sections=[('antecedentes','Antecedentes'),('justificacion','Justificación / necesidad'),('especificacion','Introducción de especificaciones'),
('calidad','Calidad',True),('ambito','Ámbito de aplicación'),('seleccion','Método de selección',True),('vigencia','Vigencia de propuesta',True),
('categoria','Categoría',True),('lugar','Lugar de entrega'),('plazo','Tiempo de entrega'),('adjudicacion','Forma de adjudicación',True),
('aceptacion','Aceptación del lote / servicio',True),('pago','Forma de pago',True),('multas','Aplicación de multas',True)]
fields += [field(*s) for s in sections]
ps=d.paragraphs
for i,p in enumerate(ps):
    t=p.text.strip()
    if t=='ADQUISICIÓN DE HERRAMIENTA PARA CUADRILLAS': replace_p(p,token('objeto'))
    elif 'BOTAS DE SEGURIDAD' in t and t.startswith('❖'): replace_p(p,token('resumen'))
    elif t=='Septiembre - 2026': replace_p(p,token('mes_anio'))
    for n,s in enumerate(sections,1):
        import re
        if re.match(r'^'+str(n)+r'\.\s+',t) and '...' not in t:
            j=i+1
            while j<len(ps) and not ps[j].text.strip(): j+=1
            if j<len(ps): replace_p(ps[j],token(s[0]))
            start=OxmlElement('w:bookmarkStart');start.set(qn('w:id'),str(100+n));start.set(qn('w:name'),f'section_{n}')
            end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),str(100+n))
            p._p.append(start);p._p.append(end)
# Index page references and running headers must belong to the new purchase.
for p in ps:
    if '...' in p.text:
        label=p.text.split('...')[0].rstrip(); n=int(label.split('.')[0])
        replace_p(p,label+'\t')
        p.paragraph_format.tab_stops.add_tab_stop(Cm(16.6),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
        word_field(p,f'PAGEREF section_{n} \\h')
seen_headers=set()
for sec in d.sections:
    for header in [sec.header,sec.first_page_header,sec.even_page_header]:
        if header.part.partname in seen_headers:continue
        seen_headers.add(header.part.partname)
        for t in header.tables:
            for row in t.rows:
                for c in row.cells:
                    for p in c.paragraphs:
                        if 'ADQUISICIÓN DE HERRAMIENTA' in p.text:replace_p(p,token('objeto'))
                        elif p.text.strip()=='SEPTIEMBRE - 2026':replace_p(p,token('mes_anio'))
                        elif '/ 7' in p.text:
                            for r in p.runs:
                                if '/ 7' in r.text:r.text=r.text.replace('/ 7','/ ')
                            word_field(p,'NUMPAGES')
update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');d.settings.element.append(update)
for c,k in zip(d.tables[0].rows[1].cells,['elaborado','revisado','aprobado']): replace_p(c.paragraphs[0],token(k))
columns=[dict(key=k,label=h.text) for k,h in zip(['numero','descripcion','unidad','cantidad','especificaciones'],d.tables[1].rows[0].cells)]
for c,k in zip(d.tables[1].rows[1].cells,[x['key'] for x in columns]): replace_p(c.paragraphs[0],token('items.'+k))
save(d,1,'Especificaciones técnicas',fields,columns,'1TDR/TDR_ENDE.docx')

# S1: native editable cells following the supplied Excel's printed layout.
import runpy
build_s1=runpy.run_path(str(ROOT/'scripts/s1-word-layout.py'))['build_s1']
d=build_s1(SRC/'2FORM S1-N014 SOL DE ADQUI/SD - FERRETERIA (1).xlsx')
rows=[('Número de solicitud','numero'),('Nombre del solicitante','solicitante'),('Cargo del solicitante','cargo'),('Área','area'),
('Verificación de existencia en almacén','almacen'),('Fecha de verificación de almacén','fecha_almacen'),('Categoría de adquisición / contratación','categoria'),
('Descripción del bien, obra o servicio','objeto'),('Posibles integrantes comité evaluador y perito externo','comite'),('Posibles proveedores','proveedores'),
('Responsable de recepción','responsable_recepcion'),('I. PREVISIÓN DE PRECIO',''),('Precio referencial (Bs)','presupuesto'),('Publicar precio','publicar_precio'),
('II. TÉRMINOS DE REFERENCIA / ESPECIFICACIONES TÉCNICAS',''),('Documento adjunto','anexo_tdr'),('Método de selección propuesto','seleccion'),('Garantías requeridas','garantias'),
('III. INFORME TÉCNICO DE JUSTIFICACIÓN DE LA NECESIDAD',''),('Justificación','justificacion'),('IV. VERIFICACIÓN PRESUPUESTARIA',''),('Con presupuesto','con_presupuesto'),('Monto verificado (Bs)','monto_verificado')]
fields=[field(k,l,k in ['categoria','seleccion','garantias']) for l,k in rows if k]+[field(k,l) for k,l in [('fecha','Fecha de solicitud'),('fecha_vobo','Fecha de visto bueno'),('fecha_aprobacion','Fecha de aprobación'),('presupuesto_usd','Precio referencial en USD'),('presupuesto_otro','Precio referencial en otra moneda')]]
save(d,2,'Solicitud de adquisición S1-N014',fields,source='2FORM S1-N014 SOL DE ADQUI/SD - FERRETERIA (1).xlsx')
catalog[-1]['version']='2026-09-10.2'

d=base('FORMULARIO DE SOLICITUD DE ADQUISICIONES Y CONTRATACIONES')
fields=[field(k,l) for k,l in [('numero','Número'),('fecha','Fecha'),('solicitante','Solicitante y cargo'),('destinatario','Destinatario y cargo'),('objeto','Objeto')]]
for f in fields: para(d,f['label']+': ',f['key'])
d.add_paragraph('CUADRO DE JUSTIFICACIÓN DE SOLICITUD DE COMPRA').runs[0].bold=True
columns=table(d,['Ítem','Cantidad','Precio unitario (Bs)','U. M.','Detalle','Activo'],['numero','cantidad','precio','unidad','descripcion','activo'])
rows=[('Tipo de compra','categoria'),('Especificaciones técnicas / ficha técnica / TDR','anexo_tdr'),('Destino final','destino'),('Centro de costo','centro_costo'),('Previsión de precio (Bs)','presupuesto'),('Justificación','justificacion'),('Lista de posibles proveedores','proveedores')]
t=d.add_table(rows=0,cols=2);t.style='Table Grid'
for l,k in rows: c=t.add_row().cells;c[0].text=l;c[1].text=token(k);fields.append(field(k,l))
signatures(d,['Solicitante','Aprobación área','Responsable de contrataciones'])
save(d,3,'Cuadro de justificación',fields,columns,'3CUADRO DE JUSTIFICACION/fotografía')

d=base('SOLICITUD DE COTIZACIÓN A EMPRESAS')
fields=[field(k,l) for k,l in [('destinatario','Empresa destinataria'),('correo','Correo destinatario'),('asunto','Asunto'),('cuerpo','Texto de la solicitud'),('fecha_limite','Fecha límite para responder'),('anexos','Documentos adjuntos'),('remitente','Remitente y cargo')]]
for f in fields:
    if f['key']=='cuerpo':
        p=d.add_paragraph(token('cuerpo'));p.paragraph_format.space_before=Pt(10);p.paragraph_format.space_after=Pt(10)
    else: para(d,f['label']+': ',f['key'])
save(d,4,'Comunicación de solicitud de cotización',fields,source='4SOL. DE COTIZACION EMPRESAS/fotografía de correo')
catalog[-1]['version']='2026-09-10.2'

d=base('SOLICITUD DE INICIO DE PROCESO')
fields=[field(k,l) for k,l in [('numero','Número'),('fecha','Fecha'),('destinatario','A'),('via','Vía'),('solicitante','De'),('objeto','Objeto'),('cuerpo','Solicitud'),('anexos','Respaldos adjuntos')]]
for f in fields: para(d,f['label']+': ',f['key'])
d.add_paragraph('Atentamente,');signatures(d,['Solicitante'])
save(d,5,'Solicitud de inicio de proceso',fields,source='5SOL DE INICIO DE PROCESO/fotografía')

d=base('SOLICITUD DE COTIZACIÓN','FORMULARIO S2-N014')
fields=[field('fecha','Fecha de solicitud'),field('destinatario','Señor(es)'),field('observaciones','Observaciones')]
para(d,'Fecha de solicitud: ','fecha');para(d,'Señor(es): ','destinatario')
d.add_paragraph('Por favor cotizar los siguientes bienes/obras/servicios:')
columns=table(d,['N.º','Cantidad','Unidad','Descripción','Precio unitario','Precio total'],['numero','cantidad','unidad','descripcion','precio_oferta','total_oferta'])
d.add_paragraph('TOTAL (Bs): ________________________')
for text in ['Tiempo de entrega: ________________________','Validez de la oferta: ________________________']: d.add_paragraph(text)
para(d,'OBSERVACIONES: ','observaciones')
d.add_paragraph('ADJUNTAR FOTOCOPIA SIMPLE DE SU RNC - NIT')
signatures(d,['Sello y firma del proveedor']);d.add_paragraph('Fecha de cotización: ________________________')
d.add_paragraph('NOTA: El presente registro no compromete una acción de compra de parte de ENDE DEORURO S.A.')
save(d,6,'Solicitud de cotización S2-N014',fields,columns,'6FORM S2-N014/fotografía')

p=SRC/'7INFORME DE CONFORMIDAD/105 -INFORME de conformidad trepaderas.docx';d=Document(p)
fields=[field(k,l) for k,l in [('fecha','Fecha'),('numero','Número del informe'),('destinatario','Destinatario'),('via','Vía'),('solicitante','Responsable informante'),('objeto','Proceso y orden de compra'),('antecedentes','Antecedentes'),('conclusiones','Conclusiones'),('recomendaciones','Recomendaciones')]]
for row,key in zip(d.tables[0].rows,[None,None,'destinatario','via','solicitante','objeto']):
    if key:
        seen=set()
        for c in row.cells[2:]:
            if id(c._tc) not in seen:
                seen.add(id(c._tc));replace_p(c.paragraphs[0],token(key))
                for p in c.paragraphs[1:]: replace_p(p,'')
seen=set()
for i,c in enumerate(d.tables[0].rows[0].cells):
    if id(c._tc) not in seen: seen.add(id(c._tc));replace_p(c.paragraphs[0],token('numero' if i==3 else 'fecha'))
ps=d.paragraphs
for i,p in enumerate(ps):
    for title,key in [('ANTECEDENTES','antecedentes'),('CONCLUSIONES','conclusiones'),('RECOMENDACIONES','recomendaciones')]:
        if p.text.strip()==title:
            j=i+1
            while j<len(ps) and not ps[j].text.strip():j+=1
            replace_p(ps[j],token(key))
    if p.text.strip()=='RECEPCIÓN DE LOS SERVICIOS':replace_p(p,'RECEPCIÓN DE BIENES / SERVICIOS')
replace_p(d.tables[1].rows[0].cells[2].paragraphs[0],'Fecha')
columns=[dict(key=k,label=h.text) for k,h in zip(['numero','descripcion','fecha_recepcion','observaciones'],d.tables[1].rows[0].cells)]
for c,col in zip(d.tables[1].rows[1].cells,columns):replace_p(c.paragraphs[0],token('items.'+col['key']))
# Do not carry embedded signatures into a new process. Keep logo media in headers.
for drawing in list(d.element.body.xpath('.//w:drawing')):
    drawing.getparent().remove(drawing)
save(d,7,'Informe de conformidad A6-N014',fields,columns,'7INFORME DE CONFORMIDAD/105 -INFORME de conformidad trepaderas.docx')

(OUT/'catalog.json').write_text(json.dumps(catalog,ensure_ascii=False,indent=2),encoding='utf-8')
print('Modelos preparados:', ', '.join(str(x['number']) for x in catalog))
