"""Read the original Excel layout and reproduce its form as native Word cells."""
from pathlib import Path
import openpyxl
from openpyxl.styles.colors import COLOR_INDEX
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.table import WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def build_s1(source):
    sheet = openpyxl.load_workbook(source).active
    doc = Document(); sec = doc.sections[0]
    sec.page_width=Cm(21); sec.page_height=Cm(29.7)
    sec.top_margin=Cm(2.7); sec.bottom_margin=Cm(2.0); sec.left_margin=Cm(2.1); sec.right_margin=Cm(2.1)
    sec.header_distance=Cm(.5); sec.footer_distance=Cm(.5)
    normal=doc.styles['Normal']; normal.font.name='Arial'; normal.font.size=Pt(8)
    normal.paragraph_format.space_after=Pt(0); normal.paragraph_format.space_before=Pt(0)
    normal.paragraph_format.line_spacing=1
    table=doc.add_table(rows=68,cols=11);table.autofit=False;table.alignment=WD_TABLE_ALIGNMENT.CENTER
    widths=[sheet.column_dimensions[openpyxl.utils.get_column_letter(c)].width if openpyxl.utils.get_column_letter(c) in sheet.column_dimensions else 8.43 for c in range(2,13)]
    widths=[(w*7+5) for w in widths]; widths=[16.8*w/sum(widths) for w in widths]
    for col,w in zip(table.columns,widths): col.width=Cm(w)
    for r,row in enumerate(table.rows,1):
        original_height=sheet.row_dimensions[r].height or sheet.sheet_format.defaultRowHeight or 15
        row.height=Pt(original_height*(.4 if original_height<=8.25 or r in [65,66,67] else .68))
        row.height_rule=WD_ROW_HEIGHT_RULE.AT_LEAST
        # Unfilled spacer cells must not impose a normal paragraph's minimum height.
        for c,cell in enumerate(row.cells,2):
            cell.width=Cm(widths[c-2]);p=cell.paragraphs[0]
            p.paragraph_format.line_spacing=Pt(1);p.add_run('').font.size=Pt(1)
            props=cell._tc.get_or_add_tcPr();margins=OxmlElement('w:tcMar')
            for edge,size in [('top',0),('bottom',0),('left',22),('right',22)]:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),str(size));e.set(qn('w:type'),'dxa');margins.append(e)
            props.append(margins)
            borders=OxmlElement('w:tcBorders'); original=sheet.cell(r,c)
            for edge in ['top','bottom','left','right']:
                side=getattr(original.border,edge);e=OxmlElement('w:'+edge)
                e.set(qn('w:val'),'single' if side.style else 'nil');e.set(qn('w:sz'),'10' if side.style in ('medium','thick') else '5');e.set(qn('w:color'),'000000');borders.append(e)
            props.append(borders)
    merges=[str(r) for r in sheet.merged_cells.ranges]
    merges += ['B10:D10','F10:L10','B12:D12','F12:L12','B14:D14','F21:L21','F42:L42','F44:L44','I8:L8','F17:H17','F19:H19','B31:L31','B38:L38','B46:L46','B51:L51','B68:E68','G68:I68','J68:L68','B33:D33','B35:C35','D35:E35','G35:H35','B53:C53','D53:F53','G53:H53','B55:C55','D55:E55']
    for area in merges:
        min_c,min_r,max_c,max_r=openpyxl.utils.range_boundaries(area)
        if min_c<2 or max_c>12:continue
        table.cell(min_r-1,min_c-2).merge(table.cell(max_r-1,max_c-2))
    replacements={
      'I8':'SOLICITUD No. {{numero}}','J8':None,'F10':'{{solicitante}}','F12':'{{cargo}}','F14':'{{area}}',
      'F17':'CON SALDO {{almacen_si}}    SIN SALDO {{almacen_no}}','H17':None,'F19':'FECHA: {{fecha_almacen}}',
      'F21':'{{categoria}}','F23':'{{objeto}}','F25':'{{comite}}','F27':'{{proveedores}}','F29':'{{responsable_recepcion}}',
      'E33':None,'F33':'{{presupuesto}}','H33':'{{presupuesto_usd}}','J33':'{{presupuesto_otro}}','B33':'PRECIO REFERENCIAL      Bs.',
      'D35':'SÍ {{publicar_precio_si}}','G35':'NO {{publicar_precio_no}}',
      'B39':'{{anexo_tdr}}','F42':'{{seleccion}}','F44':'{{garantias}}','B47':'{{justificacion}}',
      'D53':'SÍ {{con_presupuesto_si}}','F53':None,'G53':'NO {{con_presupuesto_no}}','D55':'Bs.','F55':'{{monto_verificado}}',
      'B68':'FECHA: {{fecha}}','G68':'FECHA: {{fecha_vobo}}','J68':'FECHA: {{fecha_aprobacion}}',
    }
    def color_hex(color, fallback):
        if color is None:return fallback
        if color.type=='rgb':return color.rgb[-6:]
        if color.type=='indexed' and color.indexed<len(COLOR_INDEX):return COLOR_INDEX[color.indexed][-6:]
        return fallback
    for row in sheet.iter_rows(min_row=1,max_row=68,min_col=2,max_col=12):
        for original in row:
            value=replacements.get(original.coordinate,original.value)
            if value is None:continue
            cell=table.cell(original.row-1,original.column-2);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text=str(value); p=cell.paragraphs[0];p.paragraph_format.line_spacing=1
            p.alignment={'center':WD_ALIGN_PARAGRAPH.CENTER,'right':WD_ALIGN_PARAGRAPH.RIGHT}.get(original.alignment.horizontal,WD_ALIGN_PARAGRAPH.LEFT)
            run=p.runs[0];run.font.name='Arial';run.font.size=Pt(min(9,(original.font.sz or 10)*.76));run.bold=original.font.bold
            run.font.color.rgb=RGBColor.from_string(color_hex(original.font.color,'000000'))
            if original.fill.patternType=='solid':
                fill=OxmlElement('w:shd');fill.set(qn('w:fill'),color_hex(original.fill.fgColor,'FFFFFF'));cell._tc.get_or_add_tcPr().append(fill)
            if original.coordinate in ['I8','F17','F19','B31','B38','B46','B51']:run.bold=True
            if original.coordinate in ['E33','B53']:run.font.size=Pt(7)
            if original.coordinate in ['B2','B4','B6']:
                p.alignment=WD_ALIGN_PARAGRAPH.CENTER;run.bold=True;run.font.color.rgb=RGBColor(255,255,255)
                fill=OxmlElement('w:shd');fill.set(qn('w:fill'),'003366');cell._tc.get_or_add_tcPr().append(fill)
    # Excel draws boxes across separate empty cells. Reapply their perimeter after Word merges.
    boxes=['B2:L2','B4:L4','B6:L6','I8:L8','B10:D10','F10:L10','B12:D12','F12:L12','B14:D14','F14:L14','B16:D19','F16:L19','B21:D21','F21:L21','B23:D23','F23:L23','B25:D25','F25:L25','B27:D27','F27:L27','B29:D29','F29:L29','B31:L36','B38:L40','B42:D42','F42:L42','B44:D44','F44:L44','B46:L49','B51:L57']
    for box in boxes:
        left,top,right,bottom=openpyxl.utils.range_boundaries(box)
        for r in range(top,bottom+1):
            c=2
            for rawcell in table.rows[r-1]._tr.findall(qn('w:tc')):
                props=rawcell.get_or_add_tcPr();span=props.find(qn('w:gridSpan'));end=c+(int(span.get(qn('w:val'))) if span is not None else 1)-1
                if c<left or end>right:c=end+1;continue
                borders=props.find(qn('w:tcBorders'))
                for edge,condition in [('top',r==top),('bottom',r==bottom),('left',c==left),('right',end==right)]:
                    if not condition:continue
                    node=borders.find(qn('w:'+edge))
                    if node is None:node=OxmlElement('w:'+edge);borders.append(node)
                    node.set(qn('w:val'),'single');node.set(qn('w:sz'),'10');node.set(qn('w:color'),'000000')
                c=end+1
    doc.add_paragraph().paragraph_format.line_spacing=Pt(1)
    return doc
